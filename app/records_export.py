"""Stored records as CSV files and an Excel workbook.

Every value comes from a document and is untrusted. Text a spreadsheet would run as a formula
is neutralized in both formats, and characters Excel refuses are dropped.
"""

from __future__ import annotations

import csv
import io
import math
from datetime import date, datetime, timezone
from typing import Any, Literal, NamedTuple

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ExportKind = Literal["xlsx", "records-csv", "line-items-csv"]

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"

# Leading characters that make Excel, LibreOffice, or Google Sheets read a CSV cell as a formula.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

DOCUMENT_TYPES = {"invoice": "Invoice", "receipt": "Receipt"}
PAYMENT_METHODS = {"card": "Card", "cash": "Cash", "bank": "Bank transfer"}
EXCEL_FORMATS = {"money": "#,##0.00", "date": "d mmm yyyy", "datetime": "yyyy-mm-dd hh:mm"}


class Column(NamedTuple):
    header: str
    kind: Literal["text", "number", "money", "date", "datetime"]
    width: int


RECORD_COLUMNS = (
    Column("Record ID", "number", 11),
    Column("Vendor", "text", 30),
    Column("Vendor as printed", "text", 30),
    Column("Vendor tax ID or BIN", "text", 20),
    Column("Invoice number", "text", 18),
    Column("Invoice date", "date", 13),
    Column("Due date", "date", 13),
    Column("Document type", "text", 14),
    Column("Currency", "text", 10),
    Column("Currency assumed", "text", 11),
    Column("Subtotal", "money", 14),
    Column("Tax or VAT", "money", 14),
    Column("Shipping", "money", 12),
    Column("Discount", "money", 12),
    Column("Total", "money", 15),
    Column("Payment method", "text", 15),
    Column("Checks", "text", 10),
    Column("Approved by reviewer", "text", 11),
    Column("Added (UTC)", "datetime", 17),
)

LINE_ITEM_COLUMNS = (
    Column("Record ID", "number", 11),
    Column("Vendor", "text", 30),
    Column("Invoice number", "text", 18),
    Column("Invoice date", "date", 13),
    Column("Currency", "text", 10),
    Column("Line", "number", 7),
    Column("Description", "text", 40),
    Column("Quantity", "number", 10),
    Column("Unit price", "money", 14),
    Column("Line total", "money", 14),
    Column("Category", "text", 16),
)


def _text(value: Any) -> str | None:
    return None if value is None or value == "" else str(value)


def _number(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _money(value: Any) -> float | None:
    number = _number(value)
    return None if number is None else round(number, 2)


def _day(value: Any) -> date | str | None:
    """An ISO date as a date; anything else unparseable is kept as text rather than lost."""
    text = _text(value)
    if text is None:
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        return text


def _yes_no(value: Any) -> str:
    return "Yes" if value else "No"


def _vendor(row: dict[str, Any]) -> str | None:
    """The linked vendor's chosen name, or the document's own spelling when it has no vendor."""
    return _text(row.get("vendor_name")) or _text(row["record"].get("vendor_name"))


def record_rows(rows: list[dict[str, Any]]) -> list[list[Any]]:
    """One row per record, in RECORD_COLUMNS order."""
    table = []
    for row in rows:
        record = row["record"]
        table.append([
            row["id"],
            _vendor(row),
            _text(record.get("vendor_name")),
            _text(record.get("vendor_tax_id")),
            _text(record.get("invoice_number")),
            _day(record.get("invoice_date")),
            _day(record.get("due_date")),
            DOCUMENT_TYPES.get(record.get("document_type"), _text(record.get("document_type"))),
            _text(record.get("currency")),
            _yes_no(record.get("currency_assumed")),
            _money(record.get("subtotal")),
            _money(record.get("tax_amount")),
            _money(record.get("shipping_amount")),
            _money(record.get("discount_amount")),
            _money(record.get("total_amount")),
            PAYMENT_METHODS.get(record.get("payment_method")),
            "Flagged" if row["flagged"] else "Passed",
            _yes_no(row["reviewed"]),
            row["added_at"],
        ])
    return table


def line_item_rows(rows: list[dict[str, Any]]) -> list[list[Any]]:
    """One row per line item, in LINE_ITEM_COLUMNS order, carrying its invoice's identity."""
    table = []
    for row in rows:
        record = row["record"]
        items = record.get("line_items") if isinstance(record.get("line_items"), list) else []
        for line, item in enumerate(items, 1):
            item = item if isinstance(item, dict) else {}
            table.append([
                row["id"],
                _vendor(row),
                _text(record.get("invoice_number")),
                _day(record.get("invoice_date")),
                _text(record.get("currency")),
                line,
                _text(item.get("description")),
                _number(item.get("quantity")),
                _money(item.get("unit_price")),
                _money(item.get("line_total")),
                _text(item.get("category")),
            ])
    return table


def _csv_text(text: str) -> str:
    return "'" + text if text.startswith(FORMULA_PREFIXES) else text


def _csv_value(value: Any, kind: str) -> str:
    if value is None:
        return ""
    if kind == "money":
        return f"{value:.2f}"
    if kind == "number":
        return str(int(value)) if float(value).is_integer() else str(value)
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return _csv_text(str(value))


def to_csv(columns: tuple[Column, ...], rows: list[list[Any]]) -> bytes:
    """UTF-8 with a byte-order mark, so Excel shows Bangla text and currency signs correctly."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([column.header for column in columns])
    for row in rows:
        writer.writerow([_csv_value(value, column.kind) for value, column in zip(row, columns)])
    return buffer.getvalue().encode("utf-8-sig")


def _excel_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).replace(tzinfo=None)  # Excel has no time zones
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def to_workbook(sheets: list[tuple[str, tuple[Column, ...], list[list[Any]]]]) -> bytes:
    """One sheet per table: bold frozen header, filters, column widths, and typed number and date cells."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, columns, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append([column.header for column in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([_excel_value(value) for value in row])
            for cell in sheet[sheet.max_row]:
                # openpyxl reads text starting with "=" as a formula; keep it as text.
                if cell.data_type == "f":
                    cell.data_type = "s"
        for index, column in enumerate(columns, 1):
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = column.width
            number_format = EXCEL_FORMATS.get(column.kind)
            if number_format:
                for cell in sheet[letter][1:]:
                    cell.number_format = number_format
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_export(kind: ExportKind, rows: list[dict[str, Any]], today: date) -> tuple[bytes, str, str]:
    """The file for a download: its bytes, file name, and media type."""
    stamp = today.isoformat()
    if kind == "records-csv":
        return to_csv(RECORD_COLUMNS, record_rows(rows)), f"ledgerly-records-{stamp}.csv", CSV_MEDIA_TYPE
    if kind == "line-items-csv":
        return to_csv(LINE_ITEM_COLUMNS, line_item_rows(rows)), f"ledgerly-line-items-{stamp}.csv", CSV_MEDIA_TYPE
    workbook = to_workbook([
        ("Records", RECORD_COLUMNS, record_rows(rows)),
        ("Line items", LINE_ITEM_COLUMNS, line_item_rows(rows)),
    ])
    return workbook, f"ledgerly-records-{stamp}.xlsx", XLSX_MEDIA_TYPE
