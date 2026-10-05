from __future__ import annotations

import csv
import io
from datetime import date, datetime, timedelta, timezone

from openpyxl import load_workbook

from app.records_export import (
    CSV_MEDIA_TYPE,
    LINE_ITEM_COLUMNS,
    RECORD_COLUMNS,
    XLSX_MEDIA_TYPE,
    build_export,
    line_item_rows,
    record_rows,
    to_csv,
)

ADDED = datetime(2026, 10, 4, 18, 30, tzinfo=timezone(timedelta(hours=6)))


def _row(**record: object) -> dict:
    base = {
        "vendor_name": "Tech Land BD",
        "vendor_tax_id": "001234567-0101",
        "invoice_number": "TL-1001",
        "invoice_date": "2026-03-02",
        "due_date": None,
        "document_type": "invoice",
        "currency": "BDT",
        "currency_assumed": False,
        "subtotal": 200000,
        "tax_amount": 16500.004,
        "shipping_amount": 0,
        "discount_amount": 0,
        "total_amount": 216500,
        "payment_method": "bank",
        "line_items": [
            {"description": "Laptop", "quantity": 2, "unit_price": 100000, "line_total": 200000, "category": None},
            {"description": "Mouse pad", "quantity": 0.5, "unit_price": 10, "line_total": 5, "category": "office"},
        ],
    }
    base.update(record)
    return {"id": 7, "added_at": ADDED, "record": base, "flagged": False, "reviewed": True}


def _csv_rows(content: bytes) -> list[list[str]]:
    assert content.startswith(b"\xef\xbb\xbf")  # Excel needs the byte-order mark to read UTF-8
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


def test_rows_line_up_with_their_columns() -> None:
    rows = [_row(), _row(line_items=[])]

    assert {len(row) for row in record_rows(rows)} == {len(RECORD_COLUMNS)}
    assert {len(row) for row in line_item_rows(rows)} == {len(LINE_ITEM_COLUMNS)}


def test_records_csv_spells_out_each_value() -> None:
    header, row = _csv_rows(to_csv(RECORD_COLUMNS, record_rows([_row()])))

    assert header[:4] == ["Record ID", "Vendor", "Vendor as printed", "Vendor tax ID or BIN"]
    assert dict(zip(header, row)) == {
        "Record ID": "7",
        "Vendor": "Tech Land BD",
        "Vendor as printed": "Tech Land BD",
        "Vendor tax ID or BIN": "001234567-0101",
        "Invoice number": "TL-1001",
        "Invoice date": "2026-03-02",
        "Due date": "",
        "Document type": "Invoice",
        "Currency": "BDT",
        "Currency assumed": "No",
        "Subtotal": "200000.00",
        "Tax or VAT": "16500.00",
        "Shipping": "0.00",
        "Discount": "0.00",
        "Total": "216500.00",
        "Payment method": "Bank transfer",
        "Checks": "Passed",
        "Approved by reviewer": "Yes",
        "Added (UTC)": "2026-10-04 12:30",
    }


def test_line_items_csv_repeats_the_invoice_identity() -> None:
    header, first, second = _csv_rows(to_csv(LINE_ITEM_COLUMNS, line_item_rows([_row()])))

    assert header == ["Record ID", "Vendor", "Invoice number", "Invoice date", "Currency", "Line", "Description",
                      "Quantity", "Unit price", "Line total", "Category"]
    assert first == ["7", "Tech Land BD", "TL-1001", "2026-03-02", "BDT", "1", "Laptop", "2", "100000.00",
                     "200000.00", ""]
    assert second[5:] == ["2", "Mouse pad", "0.5", "10.00", "5.00", "office"]


def test_text_that_would_run_as_a_formula_is_neutralized_in_csv() -> None:
    risky = [_row(vendor_name=text, invoice_number="-42") for text in ("=HYPERLINK(\"http://x\")", "+1", "@SUM(A1)")]

    rows = _csv_rows(to_csv(RECORD_COLUMNS, record_rows(risky)))[1:]

    assert [row[1] for row in rows] == ["'=HYPERLINK(\"http://x\")", "'+1", "'@SUM(A1)"]
    assert rows[0][4] == "'-42"
    assert rows[0][14] == "216500.00"  # amounts are numbers we format, never escaped


def test_bangla_text_and_odd_values_survive() -> None:
    row = _row(vendor_name="স্বপ্ন সুপারশপ", invoice_date="March 2", total_amount="not a number",
               document_type="cash_memo", payment_method="unknown")

    header, values = _csv_rows(to_csv(RECORD_COLUMNS, record_rows([row])))
    fields = dict(zip(header, values))

    assert fields["Vendor"] == "স্বপ্ন সুপারশপ"
    assert fields["Invoice date"] == "March 2"  # kept as written rather than dropped
    assert fields["Total"] == ""
    assert fields["Document type"] == "cash_memo"
    assert fields["Payment method"] == ""


def test_workbook_has_typed_cells_frozen_headers_and_filters() -> None:
    content, name, media_type = build_export("xlsx", [_row(), _row(line_items=[])], date(2026, 10, 5))

    workbook = load_workbook(io.BytesIO(content))
    records, lines = workbook["Records"], workbook["Line items"]
    assert (name, media_type) == ("ledgerly-records-2026-10-05.xlsx", XLSX_MEDIA_TYPE)
    assert workbook.sheetnames == ["Records", "Line items"]
    assert [cell.value for cell in records[1]][:3] == ["Record ID", "Vendor", "Vendor as printed"]
    assert records["A1"].font.bold
    assert records.freeze_panes == "A2"
    assert records.auto_filter.ref == "A1:S3"
    assert records["F2"].value == datetime(2026, 3, 2)
    assert records["F2"].number_format == "d mmm yyyy"
    assert records["O2"].value == 216500
    assert records["O2"].number_format == "#,##0.00"
    assert records["S2"].value == datetime(2026, 10, 4, 12, 30)
    assert lines.max_row == 3  # header plus the two line items of the first record


def test_workbook_keeps_formula_text_as_text_and_drops_illegal_characters() -> None:
    risky = _row(vendor_name="=cmd|' /C calc'!A0", invoice_number="INV\x07-9")

    content, _, _ = build_export("xlsx", [risky], date(2026, 10, 5))
    sheet = load_workbook(io.BytesIO(content))["Records"]

    assert sheet["B2"].value == "=cmd|' /C calc'!A0"
    assert sheet["B2"].data_type == "s"
    assert sheet["E2"].value == "INV-9"


def test_csv_exports_are_named_by_table_and_date() -> None:
    records_file = build_export("records-csv", [_row()], date(2026, 10, 5))
    lines_file = build_export("line-items-csv", [_row()], date(2026, 10, 5))

    assert records_file[1:] == ("ledgerly-records-2026-10-05.csv", CSV_MEDIA_TYPE)
    assert lines_file[1:] == ("ledgerly-line-items-2026-10-05.csv", CSV_MEDIA_TYPE)


def test_an_empty_view_still_has_headers() -> None:
    header_only = _csv_rows(build_export("line-items-csv", [], date(2026, 10, 5))[0])

    assert header_only == [[column.header for column in LINE_ITEM_COLUMNS]]


def test_vendor_column_uses_the_linked_vendor_name_and_keeps_the_printed_one() -> None:
    row = {**_row(vendor_name="RYANS"), "vendor_name": "RYANS Computers"}

    header, values = _csv_rows(to_csv(RECORD_COLUMNS, record_rows([row])))
    line = line_item_rows([row])[0]

    assert dict(zip(header, values))["Vendor"] == "RYANS Computers"
    assert dict(zip(header, values))["Vendor as printed"] == "RYANS"
    assert line[1] == "RYANS Computers"

